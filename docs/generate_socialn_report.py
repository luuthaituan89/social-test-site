from pathlib import Path
from lxml import html
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Bao_cao_do_an_SocialN.html"
OUTPUT = ROOT / "Bao_cao_do_an_SocialN.docx"
DIAGRAMS = {
    "3.1. Kiến trúc tổng thể": ("system_architecture.png", "Hình 3.1. Kiến trúc tổng thể của hệ thống SocialN"),
    "3.3. Backend": ("backend_layers.png", "Hình 3.2. Các lớp xử lý chính của backend FastAPI"),
    "4.1. Các nhóm thực thể": ("data_model.png", "Hình 4.1. Mô hình dữ liệu ER rút gọn theo miền nghiệp vụ"),
    "4.2. Luồng đăng nhập": ("auth_flow.png", "Hình 4.2. Biểu đồ tuần tự của luồng xác thực JWT"),
    "4.4. Luồng chat thời gian thực": ("chat_flow.png", "Hình 4.3. Biểu đồ tuần tự gửi và nhận tin nhắn realtime"),
    "9.5. Kiến trúc đích đề xuất": ("target_architecture.png", "Hình 9.1. Kiến trúc mục tiêu khi mở rộng SocialN lên production"),
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def plain_text(node):
    return " ".join("".join(node.itertext()).split())


def add_inline(paragraph, node):
    if node.text:
        paragraph.add_run(node.text)
    for child in node:
        text = plain_text(child)
        run = paragraph.add_run(text)
        if child.tag in ("b", "strong"):
            run.bold = True
        elif child.tag in ("i", "em"):
            run.italic = True
        elif child.tag == "code":
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        if child.tail:
            paragraph.add_run(child.tail)


def add_table(doc, node):
    rows = node.xpath("./tr|./thead/tr|./tbody/tr")
    if not rows:
        return
    columns = max(len(row.xpath("./th|./td")) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, source_row in enumerate(rows):
        cells = source_row.xpath("./th|./td")
        for c_idx, source_cell in enumerate(cells):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text = plain_text(source_cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
                    if source_cell.tag == "th" or r_idx == 0:
                        run.bold = True
            if source_cell.tag == "th" or r_idx == 0:
                set_cell_shading(cell, "E8EDF5")
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if r_idx == 0:
            set_repeat_table_header(table.rows[0])
    doc.add_paragraph()


def add_node(doc, node, context=""):
    tag = node.tag.lower() if isinstance(node.tag, str) else ""
    classes = set((node.get("class") or "").split())

    if tag == "section":
        for child in node:
            add_node(doc, child, "cover" if "cover" in classes else "toc" if "toc" in classes else context)
        doc.add_page_break()
        return
    if tag in ("h1", "h2", "h3"):
        level = int(tag[1])
        if tag == "h1" and len(doc.paragraphs) > 2 and "first" not in classes:
            doc.add_page_break()
        paragraph = doc.add_heading(plain_text(node), level=level)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        title = plain_text(node)
        if title in DIAGRAMS:
            filename, caption = DIAGRAMS[title]
            doc.add_picture(str(ROOT / "diagrams" / filename), width=Cm(15.5))
            picture_paragraph = doc.paragraphs[-1]
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.paragraph_format.space_before = Pt(8)
            picture_paragraph.paragraph_format.space_after = Pt(3)
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.first_line_indent = Cm(0)
            cap.paragraph_format.keep_with_next = True
            for run in cap.runs:
                run.italic = True
                run.font.size = Pt(10.5)
        return
    if tag == "p":
        paragraph = doc.add_paragraph()
        add_inline(paragraph, node)
        if context == "cover" or "center" in classes or "caption" in classes:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif "noindent" in classes or "note" in classes or "small" in classes:
            paragraph.paragraph_format.first_line_indent = Cm(0)
        if "school" in classes or "place" in classes or "subtitle" in classes or "title" in classes:
            for run in paragraph.runs:
                run.bold = True
        if "title" in classes:
            for run in paragraph.runs:
                run.font.size = Pt(24)
        elif "subtitle" in classes:
            for run in paragraph.runs:
                run.font.size = Pt(17)
        elif "small" in classes or "caption" in classes:
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
        if "caption" in classes:
            for run in paragraph.runs:
                run.italic = True
        return
    if tag in ("ul", "ol"):
        style = "List Bullet" if tag == "ul" else "List Number"
        for li in node.xpath("./li"):
            p = doc.add_paragraph(style=style)
            add_inline(p, li)
            p.paragraph_format.first_line_indent = Cm(0)
        return
    if tag == "table":
        add_table(doc, node)
        return
    if tag == "div" and "figure" in classes:
        return
    if tag == "pre":
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run("".join(node.itertext()).strip())
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F1F4F8")
        p._p.get_or_add_pPr().append(shading)
        return
    if tag == "hr":
        p = doc.add_paragraph()
        p.add_run("─" * 55)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    if tag == "div" and "signature" in classes:
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, child in enumerate(node.xpath("./div")[:2]):
            table.cell(0, idx).text = plain_text(child) + "\n\n\n\n\n"
            table.cell(0, idx).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    for child in node:
        add_node(doc, child, context)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    for name, size in (("Title", 24), ("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13.5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].font.all_caps = True
    for style_name in ("List Bullet", "List Number"):
        doc.styles[style_name].font.name = "Times New Roman"
        doc.styles[style_name].font.size = Pt(13)


def main():
    tree = html.fromstring(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.2)
    section.left_margin, section.right_margin = Cm(3), Cm(2.2)
    section.header_distance, section.footer_distance = Cm(1), Cm(1)
    configure_styles(doc)
    header = section.header.paragraphs[0]
    header.text = "BÁO CÁO ĐỒ ÁN SOCIALN"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"; run.font.size = Pt(9); run.font.italic = True
    add_page_number(section.footer.paragraphs[0])
    for node in tree.xpath("/html/body/*"):
        add_node(doc, node)
    doc.core_properties.title = "Báo cáo đồ án xây dựng mạng xã hội SocialN"
    doc.core_properties.subject = "Phân tích kiến trúc, chức năng và mã nguồn SocialN"
    doc.core_properties.author = "Sinh viên thực hiện"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
